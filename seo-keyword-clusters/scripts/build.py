import re
from collections import Counter, defaultdict
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

r = pd.read_pickle('res.pkl').reset_index(drop=True)
r['Keyword'] = r['Keyword'].astype(str)
r['Order'] = range(1, len(r) + 1)  # original position in the supplied list (top = head terms)
for col in ['Location', 'LocationHub', 'Industry', 'ServiceType', 'ToolBrand', 'NearMe', 'TypoFixed', 'Reason']:
    r[col] = r[col].fillna('')
r['URL'] = r['URL'].fillna('')


def slug(s):
    return re.sub(r'[^a-z0-9]+', '-', s.lower()).strip('-')


# ------------------------------------------------------------- post-level topics inside blog categories
STOP = set('''seo search engine optimization optimisation optimize optimise for the a an and or of to in on is are what
how why does do can should will when where who which with your my you i it its this that best top good free 2015 2016
2017 2018 2019 2020 2021 2022 2023 2024 2025 2026 2027 guide tips vs s new use using about from by be get make
website websites site sites google page pages me we our us more most than way ways need really still into without
does did doing done have has much many work works'''.split())


def meaningful(tokens):
    return [t for t in tokens if t not in STOP and not t.isdigit() and len(t) > 1]


GENERIC = set("""web online one st help digital business businesses marketing list impact affect affects practice practices single
content important find target report example examples type types definition strategy strategies ranking rankings rank
analysis analyze analyse analytic analytics performance optimized organic stat stats word words program industry
local technical expert friendly brand structure help increase improve better important meaning work company service
services agency tool tools course courses learn learning step one two three online free key big take long build per right add difference between differences new top michael manager design insights""".split())


def stem(t):
    if len(t) > 5 and t.endswith('ing'):
        t = t[:-3]
    if len(t) > 4 and t.endswith('ies'):
        return t[:-3] + 'y'
    if len(t) > 3 and t.endswith('s') and not t.endswith(('ss', 'is', 'us')):
        return t[:-1]
    return t


SPLIT_SECTIONS = lambda p: (p.startswith('Blog: ') and p not in ('Blog: SEO Tool Reviews & Comparisons', 'Blog: Industry SEO Guides', 'Blog: Local Market Guides')) or p in ('Learn SEO (Courses & Training)', 'SEO Glossary')
MIN_POST = 8
r['PostTopic'] = ''
for pillar, grp in r[r.Pillar.map(lambda p: isinstance(p, str) and SPLIT_SECTIONS(p))].groupby('Pillar'):
    raw = {i: meaningful(re.findall(r'[a-z0-9]+', c)) for i, c in zip(grp.index, grp.Clean)}
    keyed, surface = {}, defaultdict(Counter)
    for i, tl in raw.items():
        ks = []
        for j in range(len(tl) - 1):
            a, b = tl[j], tl[j + 1]
            k = ' '.join(sorted((stem(a), stem(b))))
            if stem(a) != stem(b):
                ks.append(k)
                surface[k][a + ' ' + b] += 1
        for t in tl:
            k = stem(t)
            ks.append(k)
            surface[k][t] += 1
        keyed[i] = ks
    cnt = Counter()
    for ks in keyed.values():
        cnt.update(set(ks))
    cat_words = {stem(w) for w in re.findall(r'[a-z]+', pillar.lower())}
    def ok(k, n):
        parts = k.split()
        if n < MIN_POST:
            return False
        if len(parts) == 1:
            return k not in cat_words and k not in GENERIC and parts[0] not in {stem(g) for g in GENERIC} and len(k) > 2
        return not all(p in cat_words or p in GENERIC or p in {stem(g) for g in GENERIC} for p in parts)
    cands = sorted([(k, n) for k, n in cnt.items() if ok(k, n)], key=lambda x: (-len(x[0].split()), -x[1]))
    for i, ks in keyed.items():
        s_ = set(ks)
        for k, n in cands:
            if k in s_:
                r.at[i, 'PostTopic'] = surface[k].most_common(1)[0][0].title()
                break

# page key and final URL
def page_url(row):
    if row.Relevance != 'Relevant':
        return ''
    if row.PostTopic:
        t = slug(row.PostTopic) if 'seo' in row.PostTopic.lower() else 'seo-' + slug(row.PostTopic)
        return f"/blog/{t}/" if not row.Pillar.startswith(('Learn', 'SEO Glossary')) else \
            (f"/learn-seo/{t}/" if row.Pillar.startswith('Learn') else f"/glossary/{slug(row.PostTopic)}/")
    return row.URL


r['PageURL'] = r.apply(page_url, axis=1)


def page_name(row):
    if row.Relevance != 'Relevant':
        return row.Cluster
    if row.PostTopic:
        return f"{row.Cluster} › {row.PostTopic}"
    return row.Cluster


r['PageName'] = r.apply(page_name, axis=1)


def page_type(row):
    if row.PostTopic:
        return {'Learn SEO (Courses & Training)': 'Learning page', 'SEO Glossary': 'Glossary entry'}.get(row.Pillar, 'Blog post')
    if row.Pillar.startswith('Blog: ') and row.Relevance == 'Relevant' and row.Pillar != 'Blog: SEO Tool Reviews & Comparisons':
        return 'Blog category hub'
    return row.PageType


r['PageType2'] = r.apply(page_type, axis=1)


def parent(row):
    p = row.Pillar
    if row.Relevance != 'Relevant':
        return ''
    if p == 'Location Pages':
        return '/locations/ › ' + (row.LocationHub or 'Other')
    if p == 'Industry Pages':
        return '/industries/'
    if p == 'Service Pages':
        return '/seo-services/'
    if p == 'Free SEO Tools':
        return '/tools/'
    if p == 'Core Money Pages':
        return 'Main navigation'
    if row.PostTopic and p.startswith('Blog: '):
        return f"/blog/{slug(p[6:])}/ (category)"
    if p.startswith('Blog'):
        return '/blog/'
    if p.startswith('Learn'):
        return '/learn-seo/'
    if p == 'SEO Glossary':
        return '/glossary/'
    return ''


r['Parent'] = r.apply(parent, axis=1)
r['PageID'] = r.apply(lambda x: (x.PageURL or '') + ' | ' + x.PageName if x.Relevance == 'Relevant' else '', axis=1)

notes = []
for _, x in r.iterrows():
    n = []
    if x.Clean != x.Keyword.lower().strip() and re.match(r'^\s*\d', x.Keyword):
        n.append('numeric prefix removed')
    if x.TypoFixed:
        n.append('typo normalised for matching')
    if x.NearMe:
        n.append('near-me intent')
    notes.append('; '.join(n))
r['Notes'] = notes
r.to_pickle('final.pkl')

# ------------------------------------------------------------- workbook
FONT = 'Arial'
HDR_FILL = PatternFill('solid', fgColor='1F3864')
HDR_FONT = Font(name=FONT, bold=True, color='FFFFFF', size=10)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True, color='1F3864')
thin = Side(style='thin', color='D9D9D9')
SECTION_FILL = {
    'Core Money Pages': 'FCE4D6', 'Service Pages': 'FFF2CC', 'Industry Pages': 'E2EFDA', 'Location Pages': 'DDEBF7',
    'Free SEO Tools': 'EDE7F6', 'Learn SEO (Courses & Training)': 'E7F4F4', 'SEO Glossary': 'F2F2F2',
}

wb = Workbook()

# ---- Read Me
ws = wb.active
ws.title = 'Read Me'
lines = [
    ('SEO Industry Keywords – Clusters & Site Architecture', TITLE),
    ('', None),
    ('What this workbook contains', BOLD),
    ('Every keyword from the supplied list (129,742 rows) was assigned to a website section, a target page and a suggested URL, so the clusters can be used to plan the site architecture.', BODY),
    ('', None),
    ('Sheets', BOLD),
    ('Summary – keyword and page counts per website section (formulas over the Keyword Clusters sheet).', BODY),
    ('Site Architecture – one row per page to build: section, page, page type, URL, parent hub, primary keyword, keyword count and example keywords.', BODY),
    ('Keyword Clusters – every keyword with its relevance, section, cluster/page, page type, URL, search intent and detected location / industry / service / tool brand.', BODY),
    ('Irrelevant & Review – the keywords marked Irrelevant or Review, with the reason.', BODY),
    ('', None),
    ('Website sections (pillars)', BOLD),
    ('Core Money Pages – head commercial terms: SEO agency/company, SEO services, SEO consultant, SEO expert/specialist, SEO freelancer, homepage.', BODY),
    ('Service Pages – one landing page per SEO service (Local, Technical, Ecommerce, AI/GEO/AEO, Link Building, White Label, Audit, …) plus a Pricing & Packages page.', BODY),
    ('Industry Pages – one landing page per vertical (Dental, Law Firms, Home Services, Healthcare, SaaS, …) under /industries/.', BODY),
    ('Location Pages – one landing page per city / state / country found in commercial keywords, grouped under a parent hub (US state or country).', BODY),
    ('Free SEO Tools – tool pages (SEO checker, rank tracker, keyword tool, meta tag generator, …).', BODY),
    ('Blog sections – informational keywords grouped into a blog category, then split into individual post topics (Site Architecture lists each post).', BODY),
    ('Blog: SEO Tool Reviews & Comparisons – third-party tool reviews (Ahrefs, Semrush, Moz, Yoast…), alternatives and “best X tools” roundups.', BODY),
    ('Learn SEO / SEO Glossary – course & training pages, and definition (“what is X in SEO”) pages.', BODY),
    ('', None),
    ('Relevance labels', BOLD),
    ('Relevant – target it on the page shown.', BODY),
    ('Irrelevant – do not target. Reasons: non-English language, non-Latin script, Korean actors/celebrities named “Seo”, adult/brand-unsafe, spam or other-site navigational, piracy.', BODY),
    ('Review – mostly unknown words: likely competitor agency brands, people’s names, very small places or misspellings. Check them by hand.', BODY),
    ('', None),
    ('Method & assumptions', BOLD),
    ('• Rule-based classification using SEO-specific dictionaries (services, industries, tools, blog topics) and a place-name list (GeoNames cities ≥15k population, US states/counties, USPS US towns).', BODY),
    ('• Priority when a keyword fits several buckets: Irrelevant → Careers/Learning → Tool brand → Location (commercial) → Pricing → Tools → Industry → Service → Core page → Blog.', BODY),
    ('• Numeric prefixes such as “11505: ” were removed before classifying, and common misspellings (e.g. “seo agnecy”) were normalised. The original keyword is kept unchanged in column A.', BODY),
    ('• “Primary keyword” is the first keyword of each page in the original file order. The source list looks sorted by importance (it starts with seo, seo tools, seo services), but it has no search volumes. Add volume/KD data to choose primary keywords properly.', BODY),
    ('• Blog post topics are built from the most frequent shared phrase (8+ keywords) in each category. Keywords without such a phrase stay on the category hub page.', BODY),
    ('• Automated grouping of ~130k keywords is not perfect. Spot-check the large clusters (especially “General SEO – long tail (review)”) before building pages.', BODY),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t)
    if f:
        c.font = f
    c.alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 150

# ---- Keyword Clusters
kc = wb.create_sheet('Keyword Clusters')
cols = [('Keyword', 'Keyword', 42), ('Relevance', 'Relevance', 11), ('Website Section', 'Pillar', 34),
        ('Keyword Cluster', 'Cluster', 40), ('Target Page', 'PageName', 52), ('Page Type', 'PageType2', 22),
        ('Suggested URL', 'PageURL', 40), ('Parent Hub', 'Parent', 34), ('Search Intent', 'Intent', 22),
        ('Location', 'Location', 22), ('Location Hub', 'LocationHub', 28), ('Industry', 'Industry', 26),
        ('Service Type', 'ServiceType', 26), ('Tool Brand', 'ToolBrand', 18), ('Irrelevance Reason', 'Reason', 30),
        ('Notes', 'Notes', 26), ('Page ID', 'PageID', 30), ('Original Row', 'Order', 10)]
r_sorted = r.copy()
r_sorted['RelRank'] = r_sorted.Relevance.map({'Relevant': 0, 'Review': 1, 'Irrelevant': 2})
r_sorted = r_sorted.sort_values(['RelRank', 'Pillar', 'Cluster', 'PageName', 'Order'])
kc.append([c[0] for c in cols])
for row in r_sorted[[c[1] for c in cols]].itertuples(index=False):
    kc.append([('' if (v is None or (isinstance(v, float) and pd.isna(v))) else v) for v in row])
for j, (_, _, w) in enumerate(cols, 1):
    kc.column_dimensions[get_column_letter(j)].width = w
for c in kc[1]:
    c.font, c.fill = HDR_FONT, HDR_FILL
    c.alignment = Alignment(vertical='center', wrap_text=True)
kc.freeze_panes = 'B2'
NK = kc.max_row
kc.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{NK}'
for row in kc.iter_rows(min_row=2, max_row=NK):
    for c in row:
        c.font = BODY
red = Font(name=FONT, size=10, color='C00000')
amber = Font(name=FONT, size=10, color='B26B00')
for i in range(2, NK + 1):
    v = kc.cell(i, 2).value
    if v == 'Irrelevant':
        kc.cell(i, 2).font = red
    elif v == 'Review':
        kc.cell(i, 2).font = amber

# ---- Site Architecture
pages = r[r.Relevance == 'Relevant'].sort_values('Order').groupby('PageID', sort=False).agg(
    Section=('Pillar', 'first'), Page=('PageName', 'first'), PageType=('PageType2', 'first'),
    URL=('PageURL', 'first'), Parent=('Parent', 'first'), Intent=('Intent', 'first'),
    Primary=('Keyword', 'first'), N=('Keyword', 'size'), FirstOrder=('Order', 'first'),
    Examples=('Keyword', lambda s: ' | '.join(list(s)[1:9]))).reset_index()
SECTION_ORDER = ['Core Money Pages', 'Service Pages', 'Industry Pages', 'Location Pages', 'Free SEO Tools',
                 'Blog: SEO Tool Reviews & Comparisons', 'Learn SEO (Courses & Training)', 'SEO Glossary']
pages['SecRank'] = pages.Section.map(lambda s: SECTION_ORDER.index(s) if s in SECTION_ORDER else 20)
pages = pages.sort_values(['SecRank', 'Section', 'N'], ascending=[True, True, False])
pages['Priority'] = pages.N.map(lambda n: 'High' if n >= 100 else ('Medium' if n >= 25 else 'Low'))

sa = wb.create_sheet('Site Architecture', 1)
hdr = ['Website Section', 'Target Page', 'Page Type', 'Suggested URL', 'Parent Hub', 'Search Intent',
       'Primary Keyword', '# Keywords', 'Build Priority', 'Example Secondary Keywords', 'Page ID']
sa.append(hdr)
for i, p in enumerate(pages.itertuples(index=False), start=2):
    sa.append([p.Section, p.Page, p.PageType, p.URL, p.Parent, p.Intent, p.Primary,
               int(p.N),
               f'=IF(H{i}>=100,"High",IF(H{i}>=25,"Medium","Low"))', p.Examples, p.PageID])
NP = sa.max_row
widths = [34, 52, 22, 42, 34, 20, 36, 11, 12, 110, 30]
for j, w in enumerate(widths, 1):
    sa.column_dimensions[get_column_letter(j)].width = w
for c in sa[1]:
    c.font, c.fill = HDR_FONT, HDR_FILL
    c.alignment = Alignment(vertical='center', wrap_text=True)
for row in sa.iter_rows(min_row=2, max_row=NP):
    sec = row[0].value
    fill = PatternFill('solid', fgColor=SECTION_FILL.get(sec, 'FFFFFF' if not sec.startswith('Blog') else 'FDF6E3'))
    for c in row:
        c.font = BODY
    row[0].fill = fill
sa.freeze_panes = 'C2'
sa.auto_filter.ref = f'A1:K{NP}'
sa.column_dimensions['K'].hidden = False

# ---- Irrelevant & Review
ir = wb.create_sheet('Irrelevant & Review')
ir.append(['Keyword', 'Relevance', 'Reason / Cluster', 'Original Row'])
sub = r[r.Relevance != 'Relevant'].sort_values(['Relevance', 'Cluster', 'Order'])
for x in sub.itertuples(index=False):
    ir.append([x.Keyword, x.Relevance, x.Cluster, x.Order])
for j, w in enumerate([50, 12, 60, 12], 1):
    ir.column_dimensions[get_column_letter(j)].width = w
for c in ir[1]:
    c.font, c.fill = HDR_FONT, HDR_FILL
for row in ir.iter_rows(min_row=2):
    for c in row:
        c.font = BODY
ir.freeze_panes = 'A2'
ir.auto_filter.ref = f'A1:D{ir.max_row}'

# ---- Summary
sm = wb.create_sheet('Summary', 1)
sm['A1'] = 'Keyword Clustering Summary'
sm['A1'].font = TITLE
sm['A3'], sm['B3'], sm['C3'] = 'Relevance', '# Keywords', '% of total'
for c in sm[3][:3]:
    c.font, c.fill = HDR_FONT, HDR_FILL
rel = ['Relevant', 'Review', 'Irrelevant']
for i, v in enumerate(rel, 4):
    sm.cell(i, 1, v).font = BODY
    sm.cell(i, 2, f"=COUNTIF('Keyword Clusters'!$B$2:$B${NK},A{i})").font = BODY
    sm.cell(i, 3, f'=IF($B$7=0,0,B{i}/$B$7)').font = BODY
    sm.cell(i, 3).number_format = '0.0%'
sm.cell(7, 1, 'Total').font = BOLD
sm.cell(7, 2, '=SUM(B4:B6)').font = BOLD
sm.cell(7, 3, '=SUM(C4:C6)').font = BOLD
sm.cell(7, 3).number_format = '0.0%'

sm['A9'], sm['B9'], sm['C9'], sm['D9'] = 'Website Section', '# Keywords', '# Pages to build', '# High-priority pages'
for c in sm[9][:4]:
    c.font, c.fill = HDR_FONT, HDR_FILL
secs = [s for s in SECTION_ORDER] + sorted(s for s in r.Pillar.unique() if s not in SECTION_ORDER and s.startswith('Blog')) + \
       ['Needs manual review', 'Irrelevant']
start = 10
for i, s in enumerate(secs, start):
    sm.cell(i, 1, s).font = BODY
    sm.cell(i, 2, f"=COUNTIF('Keyword Clusters'!$C$2:$C${NK},A{i})").font = BODY
    sm.cell(i, 3, f"=COUNTIF('Site Architecture'!$A$2:$A${NP},A{i})").font = BODY
    sm.cell(i, 4, f"=COUNTIFS('Site Architecture'!$A$2:$A${NP},A{i},'Site Architecture'!$I$2:$I${NP},\"High\")").font = BODY
end = start + len(secs) - 1
sm.cell(end + 1, 1, 'Total').font = BOLD
for col in 'BCD':
    sm[f'{col}{end + 1}'] = f'=SUM({col}{start}:{col}{end})'
    sm[f'{col}{end + 1}'].font = BOLD
sm.cell(end + 4, 1, "Note: '# Keywords' on the Site Architecture sheet is a static count made when the workbook was built. To re-count a page after editing, filter Keyword Clusters by its Page ID (column Q).").font = Font(name=FONT, size=9, italic=True)
sm.cell(end + 3, 1, 'Build Priority rule: High = 100+ keywords on the page, Medium = 25–99, Low = under 25 (keyword count used as a proxy because the source file has no search volume).').font = Font(name=FONT, size=9, italic=True)
sm.column_dimensions['A'].width = 52
for col in 'BCD':
    sm.column_dimensions[col].width = 20
for row in sm.iter_rows(min_row=4, max_row=end + 1):
    for c in row[1:4]:
        if c.column in (2, 3, 4) and c.number_format == 'General':
            c.number_format = '#,##0'

wb.save('SEO_Keyword_Clusters_Site_Architecture.xlsx')
print('pages', NP - 1, 'keywords', NK - 1)
print(pages.groupby('Section').size())
